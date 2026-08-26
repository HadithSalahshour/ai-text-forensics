"""Read .txt / .md / .docx / .pdf into normalized plain text.

Normalization is deliberately minimal. Curly quotes, em-dashes, and spacing
quirks are *detection signals* — flattening them here would destroy the very
evidence Tier D looks for. We only fix things that are artifacts of the file
format itself (line-wrapping, hyphenation, CRLF).
"""

from __future__ import annotations

import re
import unicodedata
from pathlib import Path

SUPPORTED_SUFFIXES = {".txt", ".md", ".markdown", ".docx", ".pdf"}

# Markdown syntax to strip, leaving the prose behind.
_MD_PATTERNS: list[tuple[re.Pattern[str], str]] = [
    (re.compile(r"^```.*?^```", re.MULTILINE | re.DOTALL), ""),  # fenced code
    (re.compile(r"^\s{0,3}#{1,6}\s+", re.MULTILINE), ""),  # headings
    (re.compile(r"^\s{0,3}>\s?", re.MULTILINE), ""),  # blockquotes
    (re.compile(r"^\s{0,3}([-*_])\s*(?:\1\s*){2,}$", re.MULTILINE), ""),  # hr
    (re.compile(r"^\s{0,3}[-*+]\s+", re.MULTILINE), ""),  # bullets
    (re.compile(r"^\s{0,3}\d+[.)]\s+", re.MULTILINE), ""),  # ordered list
    (re.compile(r"!\[[^\]]*\]\([^)]*\)"), ""),  # images
    (re.compile(r"\[([^\]]+)\]\([^)]*\)"), r"\1"),  # links -> label
    (re.compile(r"`([^`]+)`"), r"\1"),  # inline code
    (re.compile(r"(\*\*|__)(.+?)\1", re.DOTALL), r"\2"),  # bold
    (re.compile(r"(?<!\*)\*(?!\*)([^*\n]+)\*(?!\*)"), r"\1"),  # italic
]


class UnsupportedFormatError(ValueError):
    pass


def _normalize_whitespace(text: str) -> str:
    """Fix format artifacts only. Leaves punctuation characters untouched."""
    text = unicodedata.normalize("NFC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace(" ", " ")  # nbsp -> space
    text = re.sub(r"[ \t]+\n", "\n", text)  # trailing space on lines
    text = re.sub(r"\n{3,}", "\n\n", text)  # collapse blank-line runs
    return text.strip()


def _reflow_hard_wrapped(text: str) -> str:
    """Join lines that were wrapped by the file format, not by the author.

    PDFs and some plain-text files break lines mid-sentence. A blank line still
    means a paragraph break; a single newline between two lowercase-ish
    fragments does not.
    """
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # de-hyphenate across breaks
    # A single newline is a real break only when the next line starts a new
    # sentence (capital / digit / bullet) or the previous line ended one.
    text = re.sub(r"(?<=[^\n.!?:;\"'\)\]])\n(?=[a-z,;)])", " ", text)
    return text


def _strip_markdown(text: str) -> str:
    for pattern, repl in _MD_PATTERNS:
        text = pattern.sub(repl, text)
    return text


def from_text(text: str, *, source_format: str = "txt", source_name: str = "input") -> dict:
    if source_format in {"md", "markdown"}:
        text = _strip_markdown(text)
    cleaned = _normalize_whitespace(_reflow_hard_wrapped(_normalize_whitespace(text)))
    return {"text": cleaned, "source_format": source_format, "source_name": source_name}


def _read_docx(path: Path) -> str:
    from docx import Document as DocxDocument  # imported lazily; heavy-ish

    doc = DocxDocument(str(path))
    blocks: list[str] = []
    for para in doc.paragraphs:
        content = para.text.strip()
        if content:
            blocks.append(content)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n\n".join(p.strip() for p in pages if p.strip())


def from_file(path: str | Path) -> dict:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)

    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnsupportedFormatError(
            f"{suffix!r} not supported. Use one of: {', '.join(sorted(SUPPORTED_SUFFIXES))}"
        )

    if suffix == ".docx":
        raw, fmt = _read_docx(path), "docx"
    elif suffix == ".pdf":
        raw, fmt = _read_pdf(path), "pdf"
    elif suffix in {".md", ".markdown"}:
        raw, fmt = path.read_text(encoding="utf-8", errors="replace"), "md"
    else:
        raw, fmt = path.read_text(encoding="utf-8", errors="replace"), "txt"

    return from_text(raw, source_format=fmt, source_name=path.name)
