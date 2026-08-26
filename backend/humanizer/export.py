"""Export rewritten text and score reports.

Plain text and Markdown are always available. DOCX needs python-docx (already a
core dependency) and PDF needs reportlab (the `export` extra); both degrade to a
clear error rather than a traceback when missing.
"""

from __future__ import annotations

import io
from datetime import date

FORMATS = ("txt", "md", "docx", "pdf")


class ExportError(RuntimeError):
    pass


def _paragraphs(text: str) -> list[str]:
    return [block.strip() for block in text.split("\n\n") if block.strip()]


def to_txt(text: str) -> bytes:
    return text.encode("utf-8")


def to_markdown(text: str) -> bytes:
    return ("\n\n".join(_paragraphs(text)) + "\n").encode("utf-8")


def to_docx(text: str, *, title: str | None = None) -> bytes:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:  # pragma: no cover
        raise ExportError("python-docx is not installed.") from exc

    document = DocxDocument()
    if title:
        document.add_heading(title, level=1)
    for block in _paragraphs(text):
        document.add_paragraph(block)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def to_pdf(text: str, *, title: str | None = None) -> bytes:
    try:
        from reportlab.lib.pagesizes import LETTER
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:
        raise ExportError(
            "PDF export needs reportlab. Run: uv pip install -e '.[export]'"
        ) from exc

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=LETTER,
        leftMargin=inch,
        rightMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
        title=title or "Document",
    )
    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    body.leading = 15

    flow = []
    if title:
        flow.append(Paragraph(_escape(title), styles["Title"]))
        flow.append(Spacer(1, 12))
    for block in _paragraphs(text):
        flow.append(Paragraph(_escape(block), body))
        flow.append(Spacer(1, 9))

    doc.build(flow)
    return buffer.getvalue()


def _escape(text: str) -> str:
    """reportlab's Paragraph parses a small HTML subset, so raw angle brackets
    and ampersands would be swallowed or raise."""
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace("\n", "<br/>")
    )


def export(text: str, fmt: str, *, title: str | None = None) -> bytes:
    fmt = fmt.lower().lstrip(".")
    if fmt == "txt":
        return to_txt(text)
    if fmt in {"md", "markdown"}:
        return to_markdown(text)
    if fmt == "docx":
        return to_docx(text, title=title)
    if fmt == "pdf":
        return to_pdf(text, title=title)
    raise ExportError(f"Unsupported format {fmt!r}. Use one of: {', '.join(FORMATS)}")


MEDIA_TYPES = {
    "txt": "text/plain; charset=utf-8",
    "md": "text/markdown; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}


# --------------------------------------------------------------------------
# Scorecard
# --------------------------------------------------------------------------


def scorecard_markdown(detection, *, source_name: str = "document") -> str:
    """A plain-text record of what was found and why.

    Written so the reasoning survives outside the app — the evidence strings
    are the point, not the number.
    """
    lines = [
        f"# AI-writing report — {source_name}",
        "",
        f"*Generated {date.today().isoformat()}*",
        "",
        f"## Score: {detection.score:.1f}/100 — {detection.verdict}",
        "",
        f"- Words analyzed: {detection.word_count}",
        f"- Statistical tier available: {'yes' if detection.statistical_available else 'no'}",
    ]
    if detection.fitted_used:
        lines.append(f"- Calibrated model applied (raw {detection.fitted_score:.1f})")
    if detection.note:
        lines += ["", f"> {detection.note}"]

    lines += ["", "## By tier", "", "| Tier | Score |", "|---|---|"]
    for tier, value in detection.tier_scores.items():
        lines.append(f"| {tier} | {value:.1f} |")

    lines += ["", "## Strongest signals", ""]
    for signal in detection.top_signals(12):
        lines.append(f"- **{signal.label}** ({signal.tier}) — {signal.evidence}")

    lines += [
        "",
        "---",
        "",
        "Scores are indicative. Below ~150 words the underlying statistics are "
        "too noisy for a confident verdict, and no detector — including this "
        "one — should be treated as proof of authorship.",
        "",
    ]
    return "\n".join(lines)
